#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""论文格式（最终版）：按「角色」绝对赋值字号与字体，杜绝混杂值。

规格：
  封面主标题   黑体 22pt
  封面副标题   黑体 14pt
  一级标题     黑体 16pt（三号）加粗
  二级标题     黑体 14pt（四号）加粗
  三级标题     黑体 12pt（小四）加粗
  正文         宋体 10.5pt（五号），行距 1.5，首行缩进 2 字符（21pt）
  表格文字     宋体 10.5pt（五号）—— 与正文同号（9pt 太小，已废弃）
  代码块       Consolas 9pt，单倍行距，不缩进
  页脚页码     宋体 10.5pt

为什么之前"改数字没用"：旧实现只在字号 ≥12.5pt 时才压缩，比五号小的 10pt/9.5pt 原样保留，
于是正文里 10 / 10.5 / 11 / 12.5 混在一起。本版对每个 run 绝对赋值，跑完只剩上表几种值。

用法：
    python scripts/docs/thesis_format.py --dir <目录> [--dir ...] [--dry-run]
"""
import argparse
import glob
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

HEI = "黑体"
SONG = "宋体"
TNR = "Times New Roman"
MONO = ("Consolas", "Courier New", "Courier", "Menlo", "Monaco")

SIZE_COVER_TITLE = 22
SIZE_COVER_SUB = 14
SIZE_H1, SIZE_H2, SIZE_H3 = 16, 14, 12
SIZE_BODY = 10.5
SIZE_TABLE = 10.5
SIZE_CODE = 9
SIZE_FOOTER = 10.5


def _rfonts(run):
    rpr = run._element.get_or_add_rPr()
    return rpr.get_or_add_rFonts()


def set_run(run, east_asia, size, bold=None, mono=False):
    name = MONO[0] if mono else TNR
    run.font.name = name
    rf = _rfonts(run)
    rf.set(qn("w:ascii"), name)
    rf.set(qn("w:hAnsi"), name)
    rf.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_style_font(style, east_asia, size, bold=None):
    style.font.name = TNR
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rf = rpr.get_or_add_rFonts()
    rf.set(qn("w:ascii"), TNR)
    rf.set(qn("w:hAnsi"), TNR)
    rf.set(qn("w:eastAsia"), east_asia)


def is_mono(run):
    return any(m.lower() in (run.font.name or "").lower() for m in MONO)


def add_page_number(doc):
    for section in doc.sections:
        footer = section.footer
        if "PAGE" in footer._element.xml:
            continue
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        for el, attr in (("w:fldChar", {"w:fldCharType": "begin"}),
                         ("w:instrText", None),
                         ("w:fldChar", {"w:fldCharType": "end"})):
            node = OxmlElement(el)
            if attr:
                for k, v in attr.items():
                    node.set(qn(k), v)
            else:
                node.set(qn("xml:space"), "preserve")
                node.text = " PAGE "
            run._r.append(node)
        set_run(run, SONG, SIZE_FOOTER)


def format_docx(path, dry_run=False):
    doc = Document(path)
    stats = {"cover": 0, "h": 0, "body": 0, "table": 0, "code": 0}

    styles = doc.styles
    set_style_font(styles["Normal"], SONG, SIZE_BODY)
    pf = styles["Normal"].paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = Pt(SIZE_BODY * 2)

    for name, size in (("Heading 1", SIZE_H1), ("Heading 2", SIZE_H2), ("Heading 3", SIZE_H3)):
        try:
            st = styles[name]
        except KeyError:
            continue
        set_style_font(st, HEI, size, bold=True)
        st.font.color.rgb = RGBColor(0, 0, 0)
        spf = st.paragraph_format
        spf.first_line_indent = Pt(0)
        spf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        spf.space_before = Pt(12 if name == "Heading 1" else 6)
        spf.space_after = Pt(6)

    for name in ("List Bullet", "List Number", "List Paragraph"):
        if name in [s.name for s in styles]:
            set_style_font(styles[name], SONG, SIZE_BODY)

    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style_name = p.style.name or ""
        runs = [r for r in p.runs if r.text.strip()]
        if not runs:
            continue

        if all(is_mono(r) for r in runs):
            for r in runs:
                set_run(r, SONG, SIZE_CODE, mono=True)
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            stats["code"] += 1
            continue

        if style_name.startswith("Heading"):
            last = style_name[-1]
            size = SIZE_H1 if last == "1" else (SIZE_H2 if last == "2" else SIZE_H3)
            for r in runs:
                set_run(r, HEI, size, bold=True)
            p.paragraph_format.first_line_indent = Pt(0)
            stats["h"] += 1
            continue

        max_size = max((r.font.size.pt for r in runs if r.font.size), default=0)
        if p.alignment == WD_ALIGN_PARAGRAPH.CENTER or max_size >= 18:
            size = SIZE_COVER_TITLE if max_size >= 18 else SIZE_COVER_SUB
            for r in runs:
                set_run(r, HEI, size)
            p.paragraph_format.first_line_indent = Pt(0)
            stats["cover"] += 1
            continue

        for r in runs:
            if is_mono(r):
                set_run(r, SONG, SIZE_CODE, mono=True)
            else:
                set_run(r, SONG, SIZE_BODY)
        p.paragraph_format.first_line_indent = Pt(SIZE_BODY * 2)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        stats["body"] += 1

    for table in doc.tables:
        single = len(table.rows) == 1 and len(table.columns) == 1
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    runs = [r for r in p.runs if r.text.strip()]
                    if not runs:
                        continue
                    if single or all(is_mono(r) for r in runs):
                        for r in runs:
                            set_run(r, SONG, SIZE_CODE, mono=True)
                        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                        stats["code"] += 1
                        continue
                    for r in runs:
                        set_run(r, SONG, SIZE_TABLE)
                    p.paragraph_format.first_line_indent = Pt(0)
                    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                    stats["table"] += 1

    add_page_number(doc)
    if not dry_run:
        doc.save(path)
    return stats


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--file", nargs="*", default=[])
    ap.add_argument("--dir", action="append", default=[])
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    targets = list(args.file)
    for d in args.dir:
        targets += sorted(glob.glob(os.path.join(d, "*.docx")))
    targets = [t for t in targets if not os.path.basename(t).startswith("~$")]
    if not targets:
        print("没有指定文件")
        return 2

    ok = 0
    for path in targets:
        try:
            s = format_docx(path, dry_run=args.dry_run)
            print(f"  ✓ {os.path.basename(path):<38} 封面 {s['cover']:>2}  标题 {s['h']:>3}  "
                  f"正文 {s['body']:>3}  表格 {s['table']:>4}  代码 {s['code']:>3}")
            ok += 1
        except Exception as e:  # noqa: BLE001
            print(f"  ✗ {os.path.basename(path)}: {type(e).__name__}: {e}")
    print(f"\n处理 {ok} / {len(targets)}{'（dry-run）' if args.dry_run else ''}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
