#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把文档里的彩色 emoji 状态符号统一成纯 Unicode 字符。

映射：
  ✅ U+2705 → ✓ U+2713
  ❌ U+274C → ✗ U+2717
  ✕ U+2715 → ✗ U+2717
  ⚠️ / ⚠ + U+FE0F → ⚠ U+26A0（去掉变体选择符，避免显示成彩色 emoji）
  其余 U+FE0F 变体选择符一并删除
保留：✓ ✗ ○ ◐ × （本身都是 Unicode 字符）

同时把承载这些符号的 run 字体设为「宋体 / Times New Roman」，保证与论文正文一致。

用法：
    python scripts/docs/normalize_symbols.py --dir <目录> [--dir ...]
    python scripts/docs/normalize_symbols.py --dir ... --dry-run
"""
import argparse
import glob
import os

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

SONG = "宋体"
TNR = "Times New Roman"

MAP = {
    "\u2705": "\u2713",   # ✅ → ✓
    "\u274C": "\u2717",   # ❌ → ✗
    "\u2715": "\u2717",   # ✕ → ✗
    "\u2716": "\u2717",   # ✖ → ✗
    "\u2718": "\u2717",   # ✘ → ✗
    "\u2714\uFE0F": "\u2713",
    "\u2714": "\u2713",   # ✔ → ✓
    "\u2611": "\u2713",   # ☑ → ✓
    "\uFE0F": "",          # 去掉变体选择符
}


def normalize_text(s):
    for k, v in MAP.items():
        s = s.replace(k, v)
    return s


def fix_run(run):
    """返回替换次数；符号 run 顺手把字体统一成宋体/TNR。"""
    old = run.text
    if not old:
        return 0
    new = normalize_text(old)
    if new == old:
        return 0
    run.text = new
    # 只对「纯符号 run」调字体，避免破坏正常文字
    if all(ch in " \u2713\u2717\u26A0\u25CB\u25D0\u00D7" for ch in new):
        run.font.name = TNR
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.get_or_add_rFonts()
        rfonts.set(qn("w:ascii"), TNR)
        rfonts.set(qn("w:hAnsi"), TNR)
        rfonts.set(qn("w:eastAsia"), SONG)
        if run.font.size is None:
            run.font.size = Pt(12)
    return 1


def fix_paragraph(p):
    n = 0
    for run in p.runs:
        n += fix_run(run)
    # 跨 run 的极端情况：整段仍含 emoji，则合并到首个 run
    text = p.text
    if text and any(k in text for k in MAP):
        merged = normalize_text(text)
        if p.runs:
            first = p.runs[0]
            for r in list(p.runs)[1:]:
                r._element.getparent().remove(r._element)
            first.text = merged
            n += 1
    return n


def normalize_docx(path, dry_run=False):
    doc = Document(path)
    n = 0
    for p in doc.paragraphs:
        n += fix_paragraph(p)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    n += fix_paragraph(p)
    # 页眉页脚也扫一遍
    for section in doc.sections:
        for part in (section.header, section.footer):
            if part is None:
                continue
            for p in part.paragraphs:
                n += fix_paragraph(p)
    if not dry_run and n:
        doc.save(path)
    return n


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", action="append", default=[])
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    targets = []
    for d in args.dir:
        targets += sorted(glob.glob(os.path.join(d, "*.docx")))
    targets = [t for t in targets if not os.path.basename(t).startswith("~$")]

    if not targets:
        print("未找到文档")
        return 2

    total = 0
    for path in targets:
        try:
            n = normalize_docx(path, dry_run=args.dry_run)
        except Exception as e:  # noqa: BLE001
            print(f"  ✗ {os.path.basename(path)}: {type(e).__name__}: {e}")
            continue
        total += n
        flag = "·" if n == 0 else "✓"
        print(f"  {flag} {os.path.basename(path):<40} 替换 {n:>4} 处")
    print(f"\n合计替换 {total} 处{'（dry-run，未写盘）' if args.dry_run else ''}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
