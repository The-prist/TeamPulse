#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""几份设计文档共用的 DOCX 样式工具。

统一：A4 页面、微软雅黑（含 eastAsia）、标题层级、表格样式与列宽、代码块样式。
"""
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

EA = qn("w:eastAsia")
FONT_CN = "微软雅黑"
FONT_MONO = "Consolas"

# A4 内容宽度（去掉左右各 1.0 英寸边距后约 6.27 英寸）
CONTENT_WIDTH = 6.2


def _style_run(run, size=None, bold=None, mono=False, color=None):
    run.font.name = FONT_MONO if mono else FONT_CN
    if run._element.rPr is None:
        run._element.get_or_add_rPr()
    rfonts = run._element.rPr.get_or_add_rFonts()
    rfonts.set(EA, FONT_MONO if mono else FONT_CN)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def new_document(title: str, subtitle: str, meta_lines: list[str]) -> Document:
    doc = Document()

    sec = doc.sections[0]
    sec.page_width = Inches(8.27)
    sec.page_height = Inches(11.69)
    sec.left_margin = Inches(1.0)
    sec.right_margin = Inches(1.0)
    sec.top_margin = Inches(1.0)
    sec.bottom_margin = Inches(1.0)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_CN
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(EA, FONT_CN)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size in (("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)):
        st = doc.styles[style_name]
        st.font.name = FONT_CN
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
        if st.element.rPr is not None and st.element.rPr.rFonts is not None:
            st.element.rPr.rFonts.set(EA, FONT_CN)

    # ---- 封面 ----
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _style_run(p.add_run(title), size=30, bold=True, color=(0x1F, 0x38, 0x64))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _style_run(p.add_run(subtitle), size=13, color=(0x59, 0x59, 0x59))

    doc.add_paragraph()
    for line in meta_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _style_run(p.add_run(line), size=10.5, color=(0x40, 0x40, 0x40))

    doc.add_page_break()
    return doc


def h1(doc, text):
    doc.add_heading(text, level=1)


def h2(doc, text):
    doc.add_heading(text, level=2)


def h3(doc, text):
    doc.add_heading(text, level=3)


def para(doc, text, size=10.5, bold=False, align=None, space_after=6):
    p = doc.add_paragraph()
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _style_run(p.add_run(text), size=size, bold=bold)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def bullet(doc, text, size=10.5):
    p = doc.add_paragraph(style="List Bullet")
    _style_run(p.add_run(text), size=size)
    p.paragraph_format.space_after = Pt(2)
    return p


def numbered(doc, text, size=10.5):
    p = doc.add_paragraph(style="List Number")
    _style_run(p.add_run(text), size=size)
    p.paragraph_format.space_after = Pt(2)
    return p


def code_block(doc, text, size=8.5):
    """灰底等宽代码块（用单格表格模拟底纹）。"""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(CONTENT_WIDTH)
    shade = OxmlElement("w:shd")
    shade.set(qn("w:val"), "clear")
    shade.set(qn("w:fill"), "F2F2F2")
    cell._tc.get_or_add_tcPr().append(shade)

    cell.paragraphs[0].text = ""
    first = True
    for line in text.split("\n"):
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        _style_run(p.add_run(line if line else " "), size=size, mono=True)
    doc.add_paragraph()


def table(doc, headers, rows, widths=None, font_size=8.5, header_fill="DCE6F1"):
    """标准数据表：表头灰蓝底加粗，正文小字，固定列宽。"""
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    if widths:
        total = sum(widths)
        widths = [w * CONTENT_WIDTH / total for w in widths]
    else:
        widths = [CONTENT_WIDTH / len(headers)] * len(headers)

    hdr = tbl.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].width = Inches(widths[i])
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        _style_run(p.add_run(str(h)), size=font_size, bold=True)
        shade = OxmlElement("w:shd")
        shade.set(qn("w:val"), "clear")
        shade.set(qn("w:fill"), header_fill)
        hdr[i]._tc.get_or_add_tcPr().append(shade)

    # 表格栅格同步，避免 Word 里列宽被重算
    grid = tbl._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for i, col in enumerate(grid.findall(qn("w:gridCol"))):
            if i < len(widths):
                col.set(qn("w:w"), str(int(widths[i] * 1440)))

    for row in rows:
        cells = tbl.add_row().cells
        for i, val in enumerate(row):
            cells[i].width = Inches(widths[i])
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            _style_run(p.add_run("" if val is None else str(val)), size=font_size)

    doc.add_paragraph()
    return tbl


def note(doc, text, size=9):
    """灰字小注。"""
    p = doc.add_paragraph()
    _style_run(p.add_run(text), size=size, color=(0x66, 0x66, 0x66))
    p.paragraph_format.space_after = Pt(8)
    return p
