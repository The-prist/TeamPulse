#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""统一状态符号为四态：✓ 已实现 / △ 部分实现 / ✗ 未实现 / — 决策不做。

处理 03 需求规格（及其镜像副本）里的：
  ◐ → △            部分实现
  ○ → ✗            未实现（总览表表头与状态列）
  ✗ 决策不做 → — 决策不做    图例
  ✗ 不做 / ✗ 本期不做 → — 不做 / — 本期不做   （范围边界：这些是「决策不做」，不是「未实现」）
  ✗ 不做 → — 不做   （总览表的列头）
  ✓ 部分实现 → △ 部分实现   （范围边界里 CI/CD 那行）

几何符号 △(U+25B3) 在宋体里有字形，且与 ✓/✗ 视觉重量接近；— 用长横，表示「有意排除」。
"""
import argparse
import glob
import os

from docx import Document
from docx.oxml.ns import qn

# 顺序重要：先处理长串，再处理单字
RULES = [
    ("✗ 决策不做", "— 决策不做"),
    ("✗ 本期不做", "— 本期不做"),
    ("✗ 不做", "— 不做"),
    ("○ 未实现", "✗ 未实现"),
    ("✓ 部分实现", "△ 部分实现"),
    ("◐ 部分", "△ 部分"),
    ("◐", "△"),
    ("○", "✗"),
]


def fix_text(s):
    for old, new in RULES:
        s = s.replace(old, new)
    return s


def fix_paragraph(p, counter):
    n = 0
    for run in p.runs:
        t = run.text
        if not t:
            continue
        new = fix_text(t)
        if new != t:
            run.text = new
            if all(ch in " \u2713\u2717\u25B3\u2014\u26A0" for ch in new):
                run.font.name = "Times New Roman"
                rpr = run._element.get_or_add_rPr()
                rfonts = rpr.get_or_add_rFonts()
                rfonts.set(qn("w:ascii"), "Times New Roman")
                rfonts.set(qn("w:hAnsi"), "Times New Roman")
                rfonts.set(qn("w:eastAsia"), "宋体")
            n += 1
    # 跨 run 兜底
    text = p.text
    if text and any(r[0] in text for r in RULES):
        merged = fix_text(text)
        if merged != text and p.runs:
            first = p.runs[0]
            for r in list(p.runs)[1:]:
                r._element.getparent().remove(r._element)
            first.text = merged
            n += 1
    counter[0] += n
    return n


def normalize(path, dry_run=False):
    doc = Document(path)
    counter = [0]
    for p in doc.paragraphs:
        fix_paragraph(p, counter)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    fix_paragraph(p, counter)
    for section in doc.sections:
        for part in (section.header, section.footer):
            if part is None:
                continue
            for p in part.paragraphs:
                fix_paragraph(p, counter)
    if not dry_run and counter[0]:
        doc.save(path)
    return counter[0]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", action="append", default=[])
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    targets = []
    for d in args.dir:
        targets += sorted(glob.glob(os.path.join(d, "*.docx")))
    targets = [t for t in targets if not os.path.basename(t).startswith("~$")]

    total = 0
    for path in targets:
        n = normalize(path, dry_run=args.dry_run)
        total += n
        print(f"  {'✓' if n else '·'} {os.path.basename(path):<40} 替换 {n:>3} 处")
    print(f"\n合计 {total} 处{'（dry-run）' if args.dry_run else ''}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
