"""用例图与 ER 图说明文档生成器。

用法：
    python docs/gen/make_diagrams_doc.py
产物：
    docs/TeamPulse_用例图与ER图.docx
"""
import pathlib
import sys

sys.path.insert(0, str(pathlib.Path(__file__).parent))

from docx.shared import Cm, Pt  # noqa: E402

from word_builder import WordDoc  # noqa: E402

DOCS_DIR = pathlib.Path(__file__).parent.parent
DIAG_DIR = DOCS_DIR / "diagrams"


def picture(doc, filename, width_cm, caption):
    """插入图片并居中，下方加图注。"""
    path = DIAG_DIR / filename
    if not path.exists():
        raise FileNotFoundError(f"找不到图片：{path}")
    p = doc.doc.add_paragraph()
    p.alignment = 1  # CENTER
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    cap = doc.doc.add_paragraph()
    cap.alignment = 1
    run = cap.add_run(caption)
    run.font.size = Pt(9)
    run.italic = True
    from word_builder import COLOR_MUTED, _set_cn_font
    run.font.color.rgb = COLOR_MUTED
    _set_cn_font(run, "宋体")


def build():
    doc = WordDoc(
        title="TeamPulse 用例图与 ER 图",
        subtitle="角色用例建模 · 实体关系建模",
    )

    # ---------------- 说明 ----------------
    doc.h1("图表清单")
    doc.table(
        ["图号", "图名", "对应文件", "说明"],
        [
            ["图 1", "角色用例图", "diagrams/TeamPulse-用例图.png / .svg", "4 个参与者、3 组用例、2 条 include 关系"],
            ["图 2", "ER 实体关系图", "diagrams/TeamPulse-ER图.png / .svg", "14 张表、两个业务域、16 条关系"],
        ],
        widths=[1.6, 4.0, 6.4, 4.0],
    )
    doc.p(
        "两张图均由 PlantUML 源码生成（`diagrams/usecase.puml` 与 `diagrams/er.puml`），"
        "改完源码重新渲染即可更新图片，无需手工绘图。"
    )
    doc.callout(
        "使用建议",
        "插入报告时优先使用 **SVG 矢量版**（放大不失真）；投稿系统只接受位图时用 PNG（140 DPI）。"
        "本文件内的图片为 PNG，已按 A4 页面宽度缩放。",
        fill="E8F1FF",
    )
    doc.page_break()

    # ================= 图 1 =================
    doc.h1("图 1  角色用例图")
    picture(doc, "TeamPulse-用例图.png", 13.2, "图 1  TeamPulse 角色用例图")
    doc.page_break()

    doc.h2("1.1 参与者说明")
    doc.p(
        "本系统有 4 个参与者。其中「打卡机」与「定时任务」是**外部触发源**而不是人员，"
        "它们和人员一样从系统外部发起动作，因此同样画成参与者。"
    )
    doc.table(
        ["参与者", "类型", "职责", "身份凭证"],
        [
            ["管理员", "人员", "维护人员与团队、注册打卡设备、配置考勤规则、查询全量记录、手工补卡、查看统计",
             "账号密码登录，JWT 中角色为 admin"],
            ["普通员工", "人员", "管理自己的日程/看板/日志/评论，查看个人考勤",
             "账号密码登录，JWT 中角色为 member"],
            ["打卡机（ESP32 / 模拟器）", "外部系统", "上报打卡报文、上报心跳",
             "设备序列号 + 预共享密钥的 HMAC 签名（无账号）"],
            ["定时任务", "系统内部", "每天 00:05 生成考勤日汇总；每 60 秒巡检设备在线状态", "无需凭证，由应用内调度触发"],
        ],
        widths=[3.6, 1.8, 8.0, 4.2],
    )

    doc.h2("1.2 用例清单")
    doc.table(
        ["分组", "用例", "管理员", "普通员工", "打卡机", "定时任务"],
        [
            ["通用", "登录获取访问令牌", "✓", "✓", "—", "—"],
            ["通用", "维护个人资料", "✓", "✓", "—", "—"],
            ["协作域", "用户管理（建号/分配工号/重置密码）", "✓", "—", "—", "—"],
            ["协作域", "团队与成员管理", "✓", "仅负责的团队", "—", "—"],
            ["协作域", "日程管理", "✓", "✓ 仅本人", "—", "—"],
            ["协作域", "看板协作（拖拽/指派）", "✓", "✓ 可见看板", "—", "—"],
            ["协作域", "工作日志", "✓", "✓", "—", "—"],
            ["协作域", "评论与回复", "✓", "✓", "—", "—"],
            ["协作域", "敏感词维护", "✓", "—", "—", "—"],
            ["协作域", "统计概览", "✓ 全公司", "✓ 个人视角", "—", "—"],
            ["考勤域", "设备管理（注册/停用/轮换密钥）", "✓", "—", "—", "—"],
            ["考勤域", "配置考勤规则", "✓", "—", "—", "—"],
            ["考勤域", "打卡记录与补卡", "✓", "—", "—", "—"],
            ["考勤域", "考勤看板", "✓", "—", "—", "—"],
            ["考勤域", "我的考勤", "✓", "✓ 仅本人", "—", "—"],
            ["考勤域", "上报打卡报文", "—", "—", "✓", "—"],
            ["考勤域", "校验设备身份与签名", "—", "—", "✓ 被包含", "—"],
            ["考勤域", "判定迟到/早退", "—", "—", "✓ 被包含", "—"],
            ["考勤域", "上报设备心跳", "—", "—", "✓", "—"],
            ["考勤域", "生成考勤日汇总", "—", "—", "—", "✓"],
            ["考勤域", "设备离线巡检", "—", "—", "—", "✓"],
        ],
        widths=[1.8, 6.4, 2.2, 2.6, 2.0, 2.4],
        font_size=9,
    )

    doc.h2("1.3 两条 include 关系")
    doc.p(
        "图中「上报打卡报文」通过 `<<include>>` 指向「校验设备身份与签名」与「判定迟到/早退」，"
        "表示这两步是**每次打卡都必须执行**的强制步骤，而不是可选的扩展分支。"
        "这也是本项目与普通 CRUD 系统的本质区别：打卡不是一次简单的写入。"
    )
    doc.p("一次打卡的完整处理顺序：")
    doc.numbered(
        [
            "字段完整性校验（deviceSn / employeeNo / punchTime 必填）",
            "设备是否注册、是否被禁用",
            "设备时钟漂移是否在 ±5 分钟内（超出拒收，防「昨天抓的包今天重发」）",
            "HMAC-SHA256 签名是否正确（防伪造与篡改）",
            "报文编号是否已处理（幂等）",
            "nonce 是否已用过（防重放，10 分钟窗口）",
            "工号是否存在于系统（不存在则留痕并标记无效）",
            "判定签到/签退，按考勤规则计算迟到/早退分钟数",
            "写入打卡记录，并立即重算当日汇总",
            "向设备下发回执（设备据此声光提示）",
        ]
    )

    doc.h2("1.4 关键用例规约（一）：上报打卡报文")
    doc.table(
        ["项", "内容"],
        [
            ["用例名称", "上报打卡报文"],
            ["参与者", "打卡机（ESP32 / 模拟器）"],
            ["前置条件", "设备已在系统中注册且未被禁用；设备已连接 WiFi 并完成 NTP 对时"],
            ["触发", "员工在打卡机上完成指纹/刷卡识别"],
            ["基本流程",
             "1. 设备通过 UART 读到工号与匹配分数\n"
             "2. 设备取当前时间，生成一次性随机串 nonce\n"
             "3. 拼规范化串 deviceSn|employeeNo|punchTime|nonce 并计算 HMAC 签名\n"
             "4. 组装 JSON，发布到 MQTT 主题 attendance/checkin（HTTP 兜底通道为 POST /iot/v1/punch）\n"
             "5. 服务端接收并执行 10 步校验与判定\n"
             "6. 写入 attendance_record，重算 attendance_daily\n"
             "7. 下发回执到 attendance/ack"],
            ["异常流程",
             "· 设备未注册 → 60001，拒收\n"
             "· 设备已禁用 → 60002，拒收\n"
             "· 签名错误 → 60003，拒收\n"
             "· 报文重复 → 返回上一次的处理结果（幂等）\n"
             "· nonce 已用 → 60004，拒收\n"
             "· 时钟漂移超限 → 60006，拒收并提示校准\n"
             "· 工号不存在 → 记录仍写入但状态为 INVALID，回执 61001"],
            ["后置条件", "打卡记录落库；当日汇总被更新；设备最近心跳被刷新；设备收到回执"],
            ["可靠性", "设备断网时记录进入本地离线队列（最多 20 条），重连后重新签名补发"],
        ],
        widths=[2.4, 12.4],
        font_size=9,
    )

    doc.h2("1.5 关键用例规约（二）：手工补卡")
    doc.table(
        ["项", "内容"],
        [
            ["用例名称", "手工补卡"],
            ["参与者", "管理员"],
            ["前置条件", "员工工号已存在于系统；当前已有生效的考勤规则"],
            ["触发", "员工忘记打卡或设备故障，管理员代为补录"],
            ["基本流程",
             "1. 管理员在「打卡记录」页点击「手工补卡」\n"
             "2. 填工号、选择打卡时间、选择类型（自动判定/上班/下班）、填写备注\n"
             "3. 服务端校验工号存在、时间格式合法\n"
             "4. 判定签到或签退并计算迟到/早退\n"
             "5. 以 source=MANUAL 写入记录（deviceSn 记为 MANUAL，不校验签名）\n"
             "6. 重算当日汇总"],
            ["异常流程", "· 工号不存在 → 61001，拒绝补卡\n· 该人该时刻已有记录 → 61003，拒绝重复补卡"],
            ["后置条件", "新增一条来源为「手工补卡」的记录，来源字段可与设备打卡区分，便于审计"],
            ["设计说明", "补卡是**新增**记录，不修改历史记录——打卡数据是事件流，任何修改都会破坏审计价值"],
        ],
        widths=[2.4, 12.4],
        font_size=9,
    )

    doc.page_break()

    # ================= 图 2 =================
    doc.h1("图 2  ER 实体关系图")
    picture(doc, "TeamPulse-ER图.png", 16.0, "图 2  TeamPulse 实体关系图（14 张表）")
    doc.page_break()

    doc.h2("2.1 实体清单")
    doc.table(
        ["域", "实体（表名）", "说明", "主要键"],
        [
            ["协作", "user", "账号、密码散列、工号、角色、状态", "PK user_id；UK username / email / employee_no"],
            ["协作", "team", "团队名称、描述、创建人", "PK team_id；UK team_name"],
            ["协作", "team_member", "团队成员与团队内角色", "PK id；UK (team_id, user_id)"],
            ["协作", "schedule", "日程：时间、地点、状态、优先级", "PK schedule_id"],
            ["协作", "worklog", "工作日志，可关联日程与团队", "PK worklog_id"],
            ["协作", "kanban", "看板（必须属于团队）", "PK kanban_id"],
            ["协作", "kanban_column", "看板列与排序", "PK column_id"],
            ["协作", "kanban_item", "看板卡片（必须指向日程）", "PK item_id"],
            ["协作", "comment", "评论（多态关联日程/日志，支持两级回复）", "PK comment_id"],
            ["协作", "sensitive_word", "敏感词与替换文本", "PK word_id；UK word"],
            ["考勤", "iot_device", "打卡机：序列号、密钥、状态、心跳", "PK id；UK device_sn"],
            ["考勤", "attendance_rule", "考勤规则：上下班、宽限、工作日", "PK id；UK rule_name"],
            ["考勤", "attendance_record", "打卡事件流（只增不改）", "PK id；UK message_id、(employee_no, punch_time)"],
            ["考勤", "attendance_daily", "每人每天一行汇总", "PK id；UK (user_id, work_date)"],
        ],
        widths=[1.4, 4.0, 6.6, 4.8],
        font_size=9,
    )

    doc.h2("2.2 关系一览")
    doc.table(
        ["关系", "基数", "含义"],
        [
            ["user → team", "1 : N", "用户创建团队"],
            ["user ↔ team", "M : N", "通过 team_member 关联，附带团队内角色"],
            ["user → schedule", "1 : N", "用户创建日程"],
            ["team → schedule", "1 : N（可空）", "日程归属团队；为空表示个人日程"],
            ["schedule → worklog", "1 : N（可空）", "日志可关联日程"],
            ["team → worklog", "1 : N（可空）", "日志可归属团队"],
            ["team → kanban", "1 : N", "看板属于团队"],
            ["kanban → kanban_column", "1 : N", "一个看板包含多列"],
            ["kanban_column → kanban_item", "1 : N", "卡片落在某一列"],
            ["schedule → kanban_item", "1 : N", "卡片是日程的投影"],
            ["user → kanban_item", "1 : N（可空）", "卡片承接人"],
            ["user → comment", "1 : N", "发表评论"],
            ["comment → comment", "1 : N（可空）", "reply_to 构成两级回复"],
            ["iot_device → attendance_record", "1 : N", "哪台设备上报"],
            ["user → attendance_record", "1 : N（可空）", "按工号匹配到人；未匹配时为空但保留记录"],
            ["user → attendance_daily", "1 : N", "每人每天一行汇总"],
            ["attendance_record ⇢ attendance_daily", "聚合", "由明细聚合生成（图中虚线）"],
        ],
        widths=[6.4, 3.2, 5.2],
        font_size=9,
    )

    doc.h2("2.3 三个关键建模决策")
    doc.p("**决策一：不建物理外键。**")
    doc.p(
        "表之间是逻辑关联，由业务层保证一致性。例如删除日程时会先删除它在所有看板上的卡片。"
        "原因：协作系统的删除顺序复杂（删团队要按序清理成员、看板、卡片），物理外键会让清理逻辑僵硬、"
        "报错信息难以翻译成用户能懂的提示。代价是要求代码更自律。"
    )
    doc.p("**决策二：看板卡片必须指向日程。**")
    doc.p(
        "`kanban_item.schedule_id` 不允许为空。卡片不是独立任务，而是日程在看板上的投影。"
        "这保证了「看板显示已完成」与「日程状态是 done」永远是同一份真相，不会出现两处不一致。"
    )
    doc.p("**决策三：`employee_no` 是两个域的唯一纽带。**")
    doc.p(
        "打卡机只知道工号，不知道系统用户 ID。服务端用 `user.employee_no` 反查用户，"
        "这既是协作域与考勤域联通的唯一字段，也是唯一允许为空的合法字段（非考勤人员可不填）。"
    )

    doc.h2("2.4 图中省略的字段")
    doc.p(
        "为保持可读性，ER 图只列出主键、唯一键、关联键与核心业务字段。省略的有："
    )
    doc.bullets(
        [
            "所有表的 `created_at` / `updated_at`（由 `BaseEntity` + MyBatis-Plus 自动填充）",
            "大部分表的 `deleted` 逻辑删除标记（注意 `attendance_record`、`team_member`、`kanban_column`、"
            "`kanban_item`、`attendance_rule` 这五张表没有该列）",
            "辅助字段：`iot_device.ip_address` / `remark`、`attendance_record.remark`、`user.avatar` 等",
        ]
    )
    doc.p(
        "完整定义以建表脚本为准：`backend/src/main/resources/db/migration/V1__init_collaboration.sql` 与 "
        "`V2__init_attendance_iot.sql`。"
    )

    doc.page_break()
    doc.h1("附录  如何修改并重新生成")
    doc.p("两张图都是 PlantUML 源码，改完重新渲染即可，不需要手工绘图工具。")
    doc.code(
        "# 源码\n"
        "docs/diagrams/usecase.puml      角色用例图\n"
        "docs/diagrams/er.puml           实体关系图\n"
        "\n"
        "# 本地渲染（需要 JDK 与 plantuml.jar）\n"
        "java -jar plantuml.jar -charset UTF-8 -tpng usecase.puml er.puml\n"
        "java -jar plantuml.jar -charset UTF-8 -tsvg usecase.puml er.puml\n"
        "\n"
        "# 在线编辑（最快，粘贴即预览）\n"
        "https://www.plantuml.com/plantuml/uml/",
    )
    doc.table(
        ["想改什么", "改哪里"],
        [
            ["清晰度", "`skinparam dpi 140`（越大越清晰、文件越大）"],
            ["中文字体", "`skinparam defaultFontName \"Microsoft YaHei\"`"],
            ["配色", "`skinparam usecase/entity { BackgroundColor ... BorderColor ... }`"],
            ["增删用例/实体", "在对应 `package` 内增加 `usecase` / `entity` 与关系行"],
            ["布局方向", "用例图首行 `left to right direction`；ER 图 `skinparam linetype ortho`"],
        ],
        widths=[4.0, 10.8],
    )

    out = DOCS_DIR / "TeamPulse_用例图与ER图.docx"
    doc.save(str(out))
    print(f"generated: {out}")


if __name__ == "__main__":
    build()
